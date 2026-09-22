"""
Εκτός εφαρμογής: Excel με μοναδικά διαστήματα από το απολογιστικό
(τελευταία calculation_version).

Στήλες (6 + 2):
  Αναγνωρισμένο από / έως  (= canonical `basis_label`)
  Υπερεργασία από / έως
  Υπερωρία από / έως
  Χτύπημα από / έως        (= χτύπημα έναρξης / λήξης εργασίας:
                             πρώτη πλήρης είσοδος και τελευταία πλήρης έξοδος,
                             όχι το αναγνωρισμένο ωράριο)
"""
from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.apologistic_snapshot import CALCULATION_VERSION
from app.db import cursor
from app.timekeeping import build_day_interval_projection

_CLOCK_PART = re.compile(
    r"(?<!\d)(\d{1,2}:\d{2})\s*[\u2013\u2014\-–]\s*(\d{1,2}:\d{2})\*?",
)
# Πραγματικά χτυπήματα: επιτρέπεται κενό όριο (π.χ. «–17:15» ή «09:00–»).
_PUNCH_PART = re.compile(
    r"(?<!\d)(?:(\d{1,2}:\d{2})\s*)?[\u2013\u2014\-–]\s*(?:(\d{1,2}:\d{2})\*?)?"
)
_HM = re.compile(r"^(\d{1,2}):(\d{2})$")


def _norm_hm(value: str | None) -> str | None:
    raw = str(value or "").strip()[:5]
    m = _HM.match(raw)
    if not m:
        return None
    return f"{int(m.group(1)):02d}:{int(m.group(2)):02d}"


def _to_min(hm: str) -> int:
    h, m = hm.split(":")
    return int(h) * 60 + int(m)


def _from_min(total: int) -> str:
    total %= 1440
    return f"{total // 60:02d}:{total % 60:02d}"


def _parse_schedule_parts(proposed: str) -> list[tuple[str, str]]:
    text = str(proposed or "").strip()
    if not text or text in ("—", "-"):
        return []
    parts: list[tuple[str, str]] = []
    for match in _CLOCK_PART.finditer(text):
        start = _norm_hm(match.group(1))
        end = _norm_hm(match.group(2))
        if start and end:
            parts.append((start, end))
    return parts


def _parse_punch_parts(recorded: str) -> list[tuple[str | None, str | None]]:
    """Parse punch_recorded incl. incomplete sides (UI clock icon)."""
    text = str(recorded or "").strip()
    if not text or text in ("—", "-", "–"):
        return []
    parts: list[tuple[str | None, str | None]] = []
    for chunk in re.split(r"[\n·]+", text):
        chunk = chunk.strip()
        if not chunk or chunk in ("—", "-", "–"):
            continue
        match = _PUNCH_PART.search(chunk)
        if not match:
            only = _norm_hm(chunk)
            if only:
                parts.append((only, None))
            continue
        start = _norm_hm(match.group(1))
        end = _norm_hm(match.group(2))
        if start or end:
            parts.append((start, end))
    return parts


def _opening_closing_punch(
    parts: list[tuple[str | None, str | None]],
) -> tuple[str | None, str | None]:
    """Start/end punches used as work opening and closing.

    Among many card rows, a complete pair (both sides) is the work punch.
    An orphan incomplete row is not the opening/closing punch.  If the day
    has only an incomplete row, the missing side stays empty.
    This is not the recognized/basis interval.
    """
    complete = [(start, end) for start, end in parts if start and end]
    if complete:
        return complete[0][0], complete[-1][1]
    if not parts:
        return None, None
    return parts[0][0], parts[-1][1]


def _load_days(version: str) -> list[dict]:
    with cursor(commit=False) as cur:
        cur.execute(
            """
            SELECT s.name AS store_name, r.store_id, r.week_from, r.week_to,
                   CAST(d.effective_json AS nvarchar(max)) AS payload
            FROM dbo.karta_apologistic_day d
            INNER JOIN dbo.karta_apologistic_run r ON r.id = d.run_id
            INNER JOIN dbo.karta_store_config s ON s.id = r.store_id
            WHERE r.calculation_version = ?
              AND r.status IN (N'draft', N'completed', N'submitted', N'locked')
              AND d.effective_json IS NOT NULL
            """,
            (version,),
        )
        rows = cur.fetchall()
    out: list[dict] = []
    for store_name, store_id, week_from, week_to, payload in rows:
        try:
            data = json.loads(payload)
        except json.JSONDecodeError:
            continue
        if not isinstance(data, dict):
            continue
        data["_store_name"] = store_name
        data["_store_id"] = int(store_id)
        data["_week_from"] = str(week_from)
        data["_week_to"] = str(week_to)
        out.append(data)
    return out


def _collect(
    days: list[dict],
) -> tuple[
    list[tuple[str, str]],
    list[tuple[str, str]],
    list[tuple[str, str]],
    list[tuple[str, str]],
    list[dict],
]:
    schedules: set[tuple[str, str]] = set()
    overworks: set[tuple[str, str]] = set()
    overtimes: set[tuple[str, str]] = set()
    punches: set[tuple[str, str]] = set()
    detail: list[dict] = []

    for source_day in days:
        day = dict(source_day)
        if "basis_label" not in day or "overwork_interval" not in day:
            if str(day.get("status") or "").lower() == "review":
                # Review rows have no finalized canonical timekeeping basis.
                # Keep recognized/overwork empty instead of falling back to proposed.
                day.update({
                    "basis_label": "",
                    "recognized_interval": "",
                    "overwork_interval": "",
                    "overwork_from": None,
                    "overwork_to": None,
                })
            else:
                canonical = build_day_interval_projection(day)
                day.update({
                    "basis_label": canonical.get("basis_label"),
                    "recognized_interval": canonical.get("recognized_interval"),
                    "overwork_interval": canonical.get("overwork_interval"),
                    "overwork_from": canonical.get("overwork_from"),
                    "overwork_to": canonical.get("overwork_to"),
                })
        parts = _parse_schedule_parts(str(day.get("basis_label") or ""))
        for start, end in parts:
            schedules.add((start, end))

        punch_parts = _parse_punch_parts(str(day.get("punch_recorded") or ""))
        punch_from, punch_to = _opening_closing_punch(punch_parts)
        if punch_from or punch_to:
            punches.add((punch_from or "", punch_to or ""))

        overwork_parts = _parse_schedule_parts(str(day.get("overwork_interval") or ""))
        ow_from = _norm_hm(day.get("overwork_from"))
        ow_to = _norm_hm(day.get("overwork_to"))
        ow = (
            (ow_from, ow_to) if ow_from and ow_to
            else (overwork_parts[0][0], overwork_parts[-1][1])
            if overwork_parts else None
        )
        if ow:
            overworks.add(ow)

        ot_from = _norm_hm(day.get("overtime_from"))
        ot_to = _norm_hm(day.get("overtime_to"))
        if ot_from and ot_to and int(day.get("overtime_minutes") or 0) > 0:
            overtimes.add((ot_from, ot_to))
        for seg in day.get("overtime_segments") or []:
            if not isinstance(seg, dict):
                continue
            sf = _norm_hm(seg.get("from"))
            st = _norm_hm(seg.get("to"))
            if sf and st and int(seg.get("minutes") or 0) > 0:
                overtimes.add((sf, st))

        if parts or ow or (ot_from and ot_to) or punch_parts:
            detail.append({
                "store": day.get("_store_name"),
                "store_id": day.get("_store_id"),
                "week_from": day.get("_week_from"),
                "afm": day.get("employee_afm"),
                "name": f"{day.get('eponymo') or ''} {day.get('onoma') or ''}".strip(),
                "work_date": day.get("work_date"),
                "proposed": day.get("proposed"),
                "basis_label": day.get("basis_label"),
                "punch_recorded": day.get("punch_recorded"),
                "sched_from": parts[0][0] if parts else None,
                "sched_to": parts[-1][1] if parts else None,
                "punch_from": punch_from,
                "punch_to": punch_to,
                "ow_from": ow[0] if ow else None,
                "ow_to": ow[1] if ow else None,
                "ot_from": ot_from if ot_from and ot_to else None,
                "ot_to": ot_to if ot_from and ot_to else None,
                "ow_min": int(day.get("overwork_minutes") or 0),
                "ot_min": int(day.get("overtime_minutes") or 0),
            })

    def _key(item: tuple[str, str]) -> tuple[int, int]:
        # Κενό όριο στο τέλος ώστε τα ημιτελή (π.χ. «–17:15») να μένουν ορατά.
        a = _to_min(item[0]) if item[0] else 10_000
        b = _to_min(item[1]) if item[1] else 10_000
        return a, b

    return (
        sorted(schedules, key=_key),
        sorted(overworks, key=_key),
        sorted(overtimes, key=_key),
        sorted(punches, key=_key),
        detail,
    )


def _style_header(ws, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def build_workbook(
    schedules: list[tuple[str, str]],
    overworks: list[tuple[str, str]],
    overtimes: list[tuple[str, str]],
    punches: list[tuple[str, str]],
    detail: list[dict],
    *,
    version: str,
    day_count: int,
) -> Workbook:
    wb = Workbook()
    unique = wb.active
    unique.title = "Μοναδικά διαστήματα"
    headers = [
        "Αναγνωρισμένο από",
        "Αναγνωρισμένο έως",
        "Υπερεργασία από",
        "Υπερεργασία έως",
        "Υπερωρία από",
        "Υπερωρία έως",
        "Χτύπημα από",
        "Χτύπημα έως",
    ]
    _style_header(unique, headers)
    n = max(len(schedules), len(overworks), len(overtimes), len(punches), 1)
    for i in range(n):
        row = i + 2
        if i < len(schedules):
            unique.cell(row, 1, schedules[i][0])
            unique.cell(row, 2, schedules[i][1])
        if i < len(overworks):
            unique.cell(row, 3, overworks[i][0])
            unique.cell(row, 4, overworks[i][1])
        if i < len(overtimes):
            unique.cell(row, 5, overtimes[i][0])
            unique.cell(row, 6, overtimes[i][1])
        if i < len(punches):
            unique.cell(row, 7, punches[i][0])
            unique.cell(row, 8, punches[i][1])
    for col in range(1, 9):
        unique.column_dimensions[get_column_letter(col)].width = 18

    meta = wb.create_sheet("Πηγή")
    meta["A1"] = "Πηγή"
    meta["B1"] = "Απολογιστικό (karta_apologistic_day.effective_json)"
    meta["A2"] = "Έκδοση υπολογισμού"
    meta["B2"] = version
    meta["A3"] = "Ημέρες που διαβάστηκαν"
    meta["B3"] = day_count
    meta["A4"] = "Μοναδικά αναγνωρισμένα"
    meta["B4"] = len(schedules)
    meta["A5"] = "Μοναδικές υπερεργασίες"
    meta["B5"] = len(overworks)
    meta["A6"] = "Μοναδικές υπερωρίες"
    meta["B6"] = len(overtimes)
    meta["A7"] = "Μοναδικά χτυπήματα"
    meta["B7"] = len(punches)
    meta["A8"] = "Εξαγωγή"
    meta["B8"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    meta["A10"] = "Σημείωση"
    meta["B10"] = (
        "Αναγνωρισμένο = canonical basis_label της ωρομέτρησης. "
        "Υπερεργασία = canonical overwork_interval της ωρομέτρησης. "
        "Η πρόταση (proposed) εμφανίζεται μόνο στη χωριστή στήλη Πρόταση. "
        "Υπερωρία = overtime_from/to του απολογιστικού. "
        "Χτύπημα από/έως = χτύπημα έναρξης και λήξης εργασίας: πρώτη πλήρης "
        "είσοδος και τελευταία πλήρης έξοδος στα καταγεγραμμένα χτυπήματα. "
        "Ορφανή ημιτελής γραμμή αγνοείται. Δεν είναι το αναγνωρισμένο ωράριο. "
        "Κενό όριο μόνο όταν όλη η ημέρα έχει μόνο είσοδο ή μόνο έξοδο."
    )
    meta.column_dimensions["A"].width = 28
    meta.column_dimensions["B"].width = 100

    detail_ws = wb.create_sheet("Αναλυτικά ανά ημέρα")
    detail_headers = [
        "Κατάστημα", "store_id", "Εβδομάδα από", "ΑΦΜ", "Ονοματεπώνυμο", "Ημερομηνία",
        "Πρόταση", "Αναγνωρισμένο από", "Αναγνωρισμένο έως",
        "Υπερεργασία από", "Υπερεργασία έως", "Υπερεργασία λεπτά",
        "Υπερωρία από", "Υπερωρία έως", "Υπερωρία λεπτά",
        "Χτύπημα από", "Χτύπημα έως",
    ]
    _style_header(detail_ws, detail_headers)
    for idx, row in enumerate(detail, 2):
        detail_ws.cell(idx, 1, row["store"])
        detail_ws.cell(idx, 2, row["store_id"])
        detail_ws.cell(idx, 3, row["week_from"])
        detail_ws.cell(idx, 4, row["afm"])
        detail_ws.cell(idx, 5, row["name"])
        detail_ws.cell(idx, 6, row["work_date"])
        detail_ws.cell(idx, 7, row["proposed"])
        detail_ws.cell(idx, 8, row["sched_from"])
        detail_ws.cell(idx, 9, row["sched_to"])
        detail_ws.cell(idx, 10, row["ow_from"])
        detail_ws.cell(idx, 11, row["ow_to"])
        detail_ws.cell(idx, 12, row["ow_min"] or None)
        detail_ws.cell(idx, 13, row["ot_from"])
        detail_ws.cell(idx, 14, row["ot_to"])
        detail_ws.cell(idx, 15, row["ot_min"] or None)
        detail_ws.cell(idx, 16, row["punch_from"])
        detail_ws.cell(idx, 17, row["punch_to"])
    for col in range(1, 18):
        detail_ws.column_dimensions[get_column_letter(col)].width = 16
    detail_ws.column_dimensions["A"].width = 22
    detail_ws.column_dimensions["E"].width = 28
    detail_ws.column_dimensions["G"].width = 28
    return wb


def main() -> Path:
    version = CALCULATION_VERSION
    days = _load_days(version)
    schedules, overworks, overtimes, punches, detail = _collect(days)
    wb = build_workbook(
        schedules, overworks, overtimes, punches, detail,
        version=version, day_count=len(days),
    )
    out_dir = ROOT / "data"
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_path = out_dir / f"apologistic_oraria_katalogos_{stamp}.xlsx"
    wb.save(out_path)
    print(f"OK {out_path}")
    print(f"version={version} days={len(days)}")
    print(
        f"schedules={len(schedules)} overwork={len(overworks)} "
        f"overtime={len(overtimes)} punches={len(punches)}"
    )
    return out_path


if __name__ == "__main__":
    main()
