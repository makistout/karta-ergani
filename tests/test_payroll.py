from datetime import date
from decimal import Decimal

from app.access_control import NAV_ITEMS, permission_for_path, permissions_for_role, has_permission
from app.payroll import (
    PARAMETER_CATALOG,
    annual_income_tax,
    build_payroll_report,
    default_parameter_map,
    hourly_from_contract,
    legal_hourly,
    payroll_for_employee,
    recalculate_employee_row,
    resolve_parameters,
    daily_wage_for,
)


def _empty_breakdown():
    return {"day": 0, "night": 0, "sunday_holiday": 0, "night_sunday_holiday": 0}


def test_resolve_parameters_picks_latest_min_wage():
    rows = [
        {"store_id": 0, "code": "min_monthly_salary", "value": "880", "valid_from": "2025-04-01"},
        {"store_id": 0, "code": "min_monthly_salary", "value": "920", "valid_from": "2026-04-01"},
        {"store_id": 0, "code": "night_percent", "value": "25", "valid_from": "2000-01-01"},
    ]
    before = resolve_parameters(rows, as_of=date(2026, 3, 15))
    after = resolve_parameters(rows, as_of=date(2026, 9, 28))
    assert before["min_monthly_salary"] == "880"
    assert after["min_monthly_salary"] == "920"
    assert after["night_percent"] == "25"


def test_resolve_parameters_picks_efka_ceiling_by_date():
    rows = [
        {"store_id": 0, "code": "efka_monthly_ceiling", "value": "7572.62", "valid_from": "2025-01-01"},
        {"store_id": 0, "code": "efka_monthly_ceiling", "value": "7761.94", "valid_from": "2026-01-01"},
    ]
    before = resolve_parameters(rows, as_of=date(2025, 12, 31))
    after = resolve_parameters(rows, as_of=date(2026, 1, 1))
    assert before["efka_monthly_ceiling"] == "7572.62"
    assert after["efka_monthly_ceiling"] == "7761.94"


def test_store_override_beats_company_default():
    rows = [
        {"store_id": 0, "code": "sixth_day_percent", "value": "40", "valid_from": "2000-01-01"},
        {"store_id": 17, "code": "sixth_day_percent", "value": "30", "valid_from": "2026-01-01"},
    ]
    company = resolve_parameters(rows, store_id=0, as_of=date(2026, 9, 1))
    store = resolve_parameters(rows, store_id=17, as_of=date(2026, 9, 1))
    assert company["sixth_day_percent"] == "40"
    assert store["sixth_day_percent"] == "30"


def test_employee_month_pays_salary_plus_zone_and_overwork():
    params = default_parameter_map()
    contract = {
        "salary": "1.100,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "ΥΠΑΛΛΗΛΟΣ",
        "weekly_work_days": "5",
    }
    employee = {
        "employee_afm": "111111111",
        "eponymo": "ΔΟΚΙΜΗ",
        "onoma": "Α",
        "premium_minutes": {"day": 32 * 60, "night": 8 * 60, "sunday_holiday": 0, "night_sunday_holiday": 0},
        "overwork_breakdown": {**_empty_breakdown(), "day": 3 * 60},
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    hourly, warnings = hourly_from_contract(contract, params)
    assert not warnings
    assert hourly == legal_hourly(contract, params) or hourly > 0
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    assert row["period_salary"] == 1100.0
    assert row["pay_base_hours"] is False
    assert row["base_pay"] == 11.04
    assert row["extra_pay"] == 23.76
    assert row["total"] == 1134.80
    night = next(line for line in row["lines"] if line["family"] == "Βάση" and line["zone"] == "Νύχτας")
    assert night["zone_percent"] == 25.0
    assert night["amount"] == 11.04
    assert "5,52" in night["formula"]
    assert "25" in night["formula"]
    overwork = next(line for line in row["lines"] if line["family"] == "Υπερεργασία")
    assert overwork["family_percent"] == 20.0
    assert overwork["amount"] == 23.76
    assert "1,2" in overwork["formula"] or "1,20" in overwork["formula"]
    assert row["allowances_total"] == 0.0
    assert row["children_count"] == 0
    assert row["efka_insurable"] == 1119.80
    assert row["efka_employee"] == 149.72
    assert row["efka_employer"] == 244.00
    assert row["after_efka"] == 985.08
    assert row["fmy"] == 64.37
    assert row["net"] == 920.71


def test_catalog_exposes_allowance_rates():
    params = default_parameter_map()
    assert params["marriage_percent"] == "10"
    assert params["marriage_marital_codes"] == "1,2,3"
    assert params["child_allowance_1_percent"] == "5"
    assert params["child_allowance_2_percent"] == "6"
    assert params["child_allowance_3_percent"] == "10"
    assert params["child_allowance_4_percent"] == "15"
    assert params["child_allowance_5plus_percent"] == "18"
    assert params["seniority_years_per_step"] == "3"
    assert params["seniority_percent_per_step"] == "5"
    assert params["seniority_max_steps"] == "6"
    assert params["allowance_base"] == "period"
    assert params["efka_pension_employee_percent"] == "6.67"
    assert params["efka_pension_employer_percent"] == "13.33"
    assert params["efka_auxiliary_employee_percent"] == "3"
    assert params["efka_teka_employee_percent"] == "3"
    assert params["efka_teka_employer_percent"] == "3"
    assert params["efka_lump_employee_percent"] == "0"
    assert params["efka_aux_004_employee_percent"] == "0"
    assert params["efka_dypa_employee_percent"] == "1.65"
    assert params["efka_monthly_ceiling"] == "7761.94"
    assert params["efka_premium_exempt"] == "ναι"
    assert params["tax_b2_percent_0"] == "20"
    assert params["tax_b2_percent_2"] == "16"
    assert params["tax_credit_0"] == "777"
    assert params["tax_credit_2"] == "1120"
    assert params["fmy_annual_salaries"] == "14"
    assert params["bonus_christmas_pay_month"] == "12"
    assert params["bonus_easter_pay_month"] == "4"
    assert params["leave_allowance_pay_month"] == "7"
    assert params["bonus_leave_coeff"] == "0.041666"
    assert params["bonus_payout"] == "μήνας_καταβολής"
    assert params["leave_payout"] == "μήνας_καταβολής"
    assert params["leave_days_5day"] == "20"


def test_employee_month_adds_egsse_allowances_from_ergani_fields():
    params = default_parameter_map()
    contract = {
        "salary": "1.100,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "ΥΠΑΛΛΗΛΟΣ",
        "weekly_work_days": "5",
        "arithmos_teknon": "2",
        "marital_status": "ΕΓΓΑΜΟΣ/Η",
        "prior_service": "5",
    }
    employee = {
        "employee_afm": "111111111",
        "eponymo": "ΔΟΚΙΜΗ",
        "onoma": "Α",
        "premium_minutes": {"day": 32 * 60, "night": 8 * 60, "sunday_holiday": 0, "night_sunday_holiday": 0},
        "overwork_breakdown": {**_empty_breakdown(), "day": 3 * 60},
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    assert row["period_salary"] == 1100.0
    assert row["base_pay"] == 11.04
    assert row["extra_pay"] == 23.76
    assert row["children_count"] == 2
    assert row["marital_status"] == "1"
    assert row["prior_service_years"] == 5.0
    marriage = next(line for line in row["lines"] if line["family"] == "Επίδομα γάμου")
    children = next(line for line in row["lines"] if line["family"] == "Επίδομα τέκνων")
    seniority = next(line for line in row["lines"] if line["family"] == "Επίδομα προϋπηρεσίας")
    assert marriage["family_percent"] == 10.0
    assert marriage["amount"] == 110.0
    assert children["family_percent"] == 6.0
    assert children["amount"] == 66.0
    assert seniority["family_percent"] == 5.0
    assert seniority["amount"] == 55.0
    assert row["allowances_total"] == 231.0
    assert row["total"] == 1365.80
    assert row["efka_insurable"] == 1350.80
    assert row["efka_employee"] == 180.60
    assert row["after_efka"] == 1185.20
    assert row["fmy"] == 65.35
    assert row["net"] == 1119.85
    assert "10" in marriage["formula"]


def test_allowance_rates_come_from_parameters_not_code():
    params = default_parameter_map()
    params["marriage_percent"] = "0"
    params["child_allowance_2_percent"] = "8"
    params["seniority_percent_per_step"] = "0"
    contract = {
        "salary": "1.000,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "1",
        "arithmos_teknon": "2",
        "marital_status": "1",
        "prior_service": "9",
    }
    employee = {
        "employee_afm": "111111111",
        "eponymo": "Α",
        "onoma": "Α",
        "day": 0,
        "night": 0,
        "sunday_holiday": 0,
        "night_sunday_holiday": 0,
        "overwork_breakdown": _empty_breakdown(),
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    marriage = next(line for line in row["lines"] if line["family"] == "Επίδομα γάμου")
    children = next(line for line in row["lines"] if line["family"] == "Επίδομα τέκνων")
    seniority = next(line for line in row["lines"] if line["family"] == "Επίδομα προϋπηρεσίας")
    assert marriage["amount"] == 0.0
    assert children["family_percent"] == 8.0
    assert children["amount"] == 80.0
    assert seniority["amount"] == 0.0
    assert row["total"] == 1080.0
    assert row["efka_insurable"] == 1080.0


def test_efka_rates_come_from_parameters_not_code():
    params = default_parameter_map()
    for code in (
        "efka_pension_employee_percent",
        "efka_health_kind_employee_percent",
        "efka_health_cash_employee_percent",
        "efka_auxiliary_employee_percent",
        "efka_dypa_employee_percent",
        "efka_heavy_employee_percent",
        "efka_pension_employer_percent",
        "efka_health_kind_employer_percent",
        "efka_health_cash_employer_percent",
        "efka_auxiliary_employer_percent",
        "efka_dypa_employer_percent",
        "efka_heavy_employer_percent",
    ):
        params[code] = "0"
    params["efka_pension_employee_percent"] = "10"
    params["efka_pension_employer_percent"] = "20"
    contract = {
        "salary": "1.000,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "1",
    }
    employee = {
        "employee_afm": "111111111",
        "eponymo": "Α",
        "onoma": "Α",
        "day": 0,
        "night": 0,
        "sunday_holiday": 0,
        "night_sunday_holiday": 0,
        "overwork_breakdown": _empty_breakdown(),
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    assert row["total"] == 1000.0
    assert row["efka_insurable"] == 1000.0
    assert row["efka_employee"] == 100.0
    assert row["efka_employer"] == 200.0
    assert row["after_efka"] == 900.0
    assert row["fmy"] == 45.93
    assert row["net"] == 854.07


def test_efka_can_include_premiums_when_parameter_off():
    params = default_parameter_map()
    params["efka_premium_exempt"] = "όχι"
    contract = {
        "salary": "1.100,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "ΥΠΑΛΛΗΛΟΣ",
        "weekly_work_days": "5",
    }
    employee = {
        "employee_afm": "111111111",
        "eponymo": "ΔΟΚΙΜΗ",
        "onoma": "Α",
        "premium_minutes": {"day": 32 * 60, "night": 8 * 60, "sunday_holiday": 0, "night_sunday_holiday": 0},
        "overwork_breakdown": {**_empty_breakdown(), "day": 3 * 60},
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    assert row["total"] == 1134.80
    assert row["efka_insurable"] == 1134.80
    assert row["efka_employee"] == 151.71


def test_recalculate_employee_row_changes_total_when_hours_edited():
    params = default_parameter_map()
    contract = {
        "salary": "1.100,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "ΥΠΑΛΛΗΛΟΣ",
        "weekly_work_days": "5",
    }
    employee = {
        "employee_afm": "111111111",
        "eponymo": "ΔΟΚΙΜΗ",
        "onoma": "Α",
        "premium_minutes": {"day": 32 * 60, "night": 8 * 60, "sunday_holiday": 0, "night_sunday_holiday": 0},
        "overwork_breakdown": {**_empty_breakdown(), "day": 3 * 60},
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    edited = {**row, "lines": [dict(line) for line in row["lines"]]}
    night = next(line for line in edited["lines"] if line["family"] == "Βάση" and line["zone"] == "Νύχτας")
    night["hours"] = 16
    updated = recalculate_employee_row(edited)
    assert updated["total"] == 1145.84
    assert next(line for line in updated["lines"] if line["zone"] == "Νύχτας")["amount"] == 22.08


def test_worker_gets_paid_for_base_day_hours():
    params = default_parameter_map()
    contract = {
        "salary": "",
        "hourly_wage": "6,00",
        "weekly_hours": "40",
        "characterization": "ΕΡΓΑΤΗΣ",
        "weekly_work_days": "5",
    }
    employee = {
        "employee_afm": "222222222",
        "eponymo": "ΕΡΓΑΤΗΣ",
        "onoma": "Β",
        "premium_minutes": {"day": 8 * 60, "night": 0, "sunday_holiday": 0, "night_sunday_holiday": 0},
        "overwork_breakdown": _empty_breakdown(),
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    assert row["pay_base_hours"] is True
    assert row["period_salary"] == 0.0
    assert row["total"] == 48.0
    assert row["efka_insurable"] == 48.0
    assert row["efka_employee"] == 6.41
    assert row["after_efka"] == 41.59
    assert row["fmy"] == 0.0
    assert row["net"] == 41.59


def test_part_time_worker_legal_hourly_uses_fulltime_day():
    params = default_parameter_map()
    contract = {
        "salary": "172,11",
        "hourly_wage": "6,62",
        "weekly_hours": "6",
        "fulltime_contract_weekly_hours": "40",
        "characterization": "0",
        "weekly_work_days": "5",
        "regime": "1",
    }
    legal = legal_hourly(contract, params)
    hourly, warnings = hourly_from_contract(contract, params)
    assert not warnings
    assert hourly == Decimal("6.62")
    assert legal == Decimal("5.14")
    employee = {
        "employee_afm": "139052918",
        "eponymo": "ΛΑΣΠΑ",
        "onoma": "ΑΡΙΑΔΝΗ",
        "premium_minutes": {
            "day": 3 * 60, "night": 3 * 60, "sunday_holiday": 0, "night_sunday_holiday": 0,
        },
        "overwork_breakdown": _empty_breakdown(),
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="week",
    )
    assert row["hourly_wage"] == 6.62
    assert row["legal_hourly"] == 5.14
    assert row["period_salary"] == 0.0
    assert row["total"] < 80
    day = next(line for line in row["lines"] if line["family"] == "Βάση" and line["zone"] == "Ημέρας")
    night = next(line for line in row["lines"] if line["family"] == "Βάση" and line["zone"] == "Νύχτας")
    assert day["amount"] == 19.86
    assert night["amount"] == 23.72
    assert row["total"] == 43.58


def test_legal_hourly_ignores_part_time_weekly_when_fulltime_copied():
    params = default_parameter_map()
    legal = legal_hourly(
        {
            "characterization": "ΕΡΓΑΤΗΣ",
            "weekly_hours": "6",
            "fulltime_contract_weekly_hours": "6",
            "weekly_work_days": "5",
        },
        params,
    )
    assert legal == Decimal("5.14")


def test_part_time_worker_daily_wage_scales_minimum():
    params = default_parameter_map()
    contract = {
        "hourly_wage": "",
        "salary": "",
        "weekly_hours": "6",
        "fulltime_contract_weekly_hours": "40",
        "characterization": "0",
        "weekly_work_days": "5",
    }
    daily = daily_wage_for(
        contract=contract,
        params=params,
        hourly=None,
        period_salary=Decimal("0"),
        allowances=Decimal("0"),
        include_allowances=False,
    )
    assert daily == Decimal("6.16")
    with_hourly = daily_wage_for(
        contract={**contract, "hourly_wage": "6,62"},
        params=params,
        hourly=Decimal("6.62"),
        period_salary=Decimal("0"),
        allowances=Decimal("0"),
        include_allowances=False,
    )
    assert with_hourly == Decimal("7.94")


def test_hourly_from_contract_ignores_total_weekly_hours_of_other_jobs():
    params = default_parameter_map()
    hourly, warnings = hourly_from_contract(
        {
            "salary": "172,11",
            "hourly_wage": "",
            "weekly_hours": "",
            "total_weekly_hours": "11",
            "characterization": "0",
        },
        params,
    )
    assert hourly is None
    assert warnings


def test_build_payroll_report_sums_employees():
    params = default_parameter_map()
    timekeeping = {
        "employees": [
            {
                "employee_afm": "111111111",
                "eponymo": "Α",
                "onoma": "Α",
                "day": 0,
                "night": 0,
                "sunday_holiday": 0,
                "night_sunday_holiday": 0,
                "overwork_breakdown": _empty_breakdown(),
                "overtime_40_breakdown": _empty_breakdown(),
                "overtime_60_breakdown": _empty_breakdown(),
                "overtime_120_breakdown": _empty_breakdown(),
                "partial_additional_12_breakdown": _empty_breakdown(),
                "sixth_day_breakdown": _empty_breakdown(),
                "sixth_day_above_48_breakdown": _empty_breakdown(),
                "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
            }
        ]
    }
    contracts = {
        "111111111": {
            "salary": "920,00",
            "hourly_wage": "",
            "weekly_hours": "40",
            "characterization": "1",
        }
    }
    report = build_payroll_report(timekeeping, contracts, params, period_type="month")
    assert report["counts"]["employees"] == 1
    assert report["grand_total"] == 920.0
    assert report["grand_efka_employee"] == 123.00
    assert report["grand_after_efka"] == 797.00
    assert report["grand_fmy"] == 25.33
    assert report["grand_net"] == 771.67
    assert report["grand_bonuses"] == 0.0
    assert report["calculation_version"] == "payroll-v8-contract-segments"


def test_annual_income_tax_matches_2026_example():
    params = default_parameter_map()
    result = annual_income_tax(15000, children=0, age_group="over_30", params=params)
    assert result["gross_tax"] == 1900.0
    assert result["credit"] == 717.0
    assert result["tax"] == 1183.0


def test_fmy_rates_come_from_parameters_not_code():
    params = default_parameter_map()
    for code in (
        "tax_b1_percent_0",
        "tax_b2_percent_0",
        "tax_b3_percent_0",
        "tax_b4_percent",
        "tax_b5_percent",
        "tax_b6_percent",
        "tax_credit_0",
    ):
        params[code] = "0"
    contract = {
        "salary": "2.000,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "1",
    }
    employee = {
        "employee_afm": "111111111",
        "eponymo": "Α",
        "onoma": "Α",
        "day": 0,
        "night": 0,
        "sunday_holiday": 0,
        "night_sunday_holiday": 0,
        "overwork_breakdown": _empty_breakdown(),
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }
    row = payroll_for_employee(
        employee=employee, contract=contract, params=params, period_type="month",
    )
    assert row["fmy"] == 0.0
    assert row["net"] == row["after_efka"]


def test_under_25_zeroes_first_two_brackets():
    params = default_parameter_map()
    result = annual_income_tax(18000, children=0, age_group="under_25", params=params)
    assert result["gross_tax"] == 0.0
    assert result["tax"] == 0.0


def test_payroll_paths_are_super_admin_only():
    assert permission_for_path("/ui/payroll", "GET") == "payroll.view"
    assert permission_for_path("/ui/payroll/parameters", "GET") == "payroll.view"
    assert permission_for_path("/api/payroll/calculate", "POST") == "payroll.view"
    assert permission_for_path("/api/payroll/export", "POST") == "payroll.view"
    assert permission_for_path("/api/store/1/efka-settings", "GET") == "payroll.view"
    assert permission_for_path("/api/store/1/efka-settings", "PUT") == "payroll.edit"
    assert permission_for_path("/api/payroll/parameters", "PUT") == "payroll.edit"
    assert has_permission("payroll.view", role="super_admin")
    assert has_permission("payroll.edit", role="super_admin")
    for role in ("accountant", "office_manager", "office", "admin", "backoffice_admin", "viewer"):
        assert not has_permission("payroll.view", role=role)
        assert not has_permission("payroll.edit", role=role)
    assert "payroll.view" not in permissions_for_role("accountant")
    assert "payroll.edit" not in permissions_for_role("accountant")
    assert all(item.get("label") != "Μισθοδοσία" for item in NAV_ITEMS)


def _salary_employee():
    return {
        "employee_afm": "111111111",
        "eponymo": "ΔΟΚΙΜΗ",
        "onoma": "Α",
        "premium_minutes": {"day": 0, "night": 0, "sunday_holiday": 0, "night_sunday_holiday": 0},
        "overwork_breakdown": _empty_breakdown(),
        "overtime_40_breakdown": _empty_breakdown(),
        "overtime_60_breakdown": _empty_breakdown(),
        "overtime_120_breakdown": _empty_breakdown(),
        "partial_additional_12_breakdown": _empty_breakdown(),
        "sixth_day_breakdown": _empty_breakdown(),
        "sixth_day_above_48_breakdown": _empty_breakdown(),
        "exception_sixth_day_above_48_breakdown": _empty_breakdown(),
    }


def _salary_contract(**extra):
    row = {
        "salary": "1.000,00",
        "hourly_wage": "",
        "weekly_hours": "40",
        "characterization": "1",
        "weekly_work_days": "5",
        "hire_date": "2025-01-01",
    }
    row.update(extra)
    return row


def test_september_has_no_seasonal_bonuses():
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=default_parameter_map(),
        period_type="month",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
    )
    assert row["bonuses_total"] == 0.0
    assert not any(line["line_kind"] == "bonus" for line in row["lines"])


def test_december_pays_full_christmas_bonus_with_leave_coeff():
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=default_parameter_map(),
        period_type="month",
        period_from=date(2026, 12, 1),
        period_to=date(2026, 12, 31),
    )
    gift = next(line for line in row["lines"] if line["family"] == "Δώρο Χριστουγέννων")
    assert gift["amount"] == 1041.67
    assert row["bonuses_total"] == 1041.67
    assert row["total"] == 2041.67
    assert row["fmy_bonus"] > 0
    assert row["tax_bonus_after"] > 0
    # Τα δώρα δεν ετήσιοποιούνται × 14.
    assert row["fmy"] < 400


def _declared_schedule(year, month, six_day=False):
    from calendar import monthrange
    return [{"work_date": date(year, month, day).isoformat(),
             "hour_from": "09:00", "hour_to": "15:40" if six_day else "17:00"}
            for day in range(1, monthrange(year, month)[1] + 1)
            if date(year, month, day).weekday() < (6 if six_day else 5)]


def test_christmas_prorates_for_short_employment():
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(hire_date="2026-12-13"),
        params=default_parameter_map(),
        schedule_rows=_declared_schedule(2026,12),
        period_type="month",
        period_from=date(2026, 12, 1),
        period_to=date(2026, 12, 31),
    )
    gift = next(line for line in row["lines"] if line["family"] == "Δώρο Χριστουγέννων")
    assert gift["amount"] == 83.33
    assert gift["hours"] == 19.0


def test_april_pays_easter_bonus():
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=default_parameter_map(),
        period_type="month",
        period_from=date(2026, 4, 1),
        period_to=date(2026, 4, 30),
    )
    gift = next(line for line in row["lines"] if line["family"] == "Δώρο Πάσχα")
    assert gift["amount"] == 520.83
    assert row["bonuses_total"] == 520.83


def test_july_pays_leave_allowance():
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=default_parameter_map(),
        period_type="month",
        period_from=date(2026, 7, 1),
        period_to=date(2026, 7, 31),
    )
    gift = next(line for line in row["lines"] if line["family"] == "Επίδομα αδείας")
    assert gift["amount"] == 800.0
    assert row["bonuses_total"] == 800.0


def test_bonus_rates_come_from_parameters_not_code():
    params = default_parameter_map()
    params["bonus_christmas_pay_month"] = "0"
    params["bonus_easter_pay_month"] = "0"
    params["leave_allowance_pay_month"] = "0"
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=params,
        period_type="month",
        period_from=date(2026, 12, 1),
        period_to=date(2026, 12, 31),
    )
    assert row["bonuses_total"] == 0.0


def test_monthly_payout_spreads_christmas_across_period():
    params = default_parameter_map()
    params["bonus_payout"] = "κάθε_μήνας_περιόδου"
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=params,
        period_type="month",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
    )
    gift = next(line for line in row["lines"] if line["family"] == "Δώρο Χριστουγέννων")
    assert gift["amount"] == 127.55
    assert gift["hours"] == 30.0
    assert not any(line["family"] == "Δώρο Πάσχα" for line in row["lines"])


def test_monthly_payout_leave_spreads_over_year():
    params = default_parameter_map()
    params["leave_payout"] = "κάθε_μήνας"
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=params,
        period_type="month",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
    )
    gift = next(line for line in row["lines"] if line["family"] == "Επίδομα αδείας")
    assert gift["amount"] == 67.95
    assert not any(line["family"] == "Δώρο Χριστουγέννων" for line in row["lines"])


def test_payroll_afm_key_pads_digits():
    from app.payroll import payroll_afm_key

    assert payroll_afm_key("12345678") == "012345678"
    assert payroll_afm_key("012345678") == "012345678"
    assert payroll_afm_key("") == ""


def test_build_payroll_matches_padded_contract_afm():
    params = default_parameter_map()
    report = build_payroll_report(
        {"employees": [{**_salary_employee(), "employee_afm": "12345678"}]},
        {"012345678": _salary_contract(hire_date="2026-01-15")},
        params,
        period_type="month",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
    )
    row = report["employees"][0]
    assert row["hire_date"] == "2026-01-15"
    assert not any("Χωρίς ημερομηνία πρόσληψης" in w for w in row["warnings"])


def test_payroll_uses_eteaep_by_default_and_teka_when_contract_says_so():
    params = default_parameter_map()
    eteaep = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=params,
        period_type="month",
    )
    aux = next(b for b in eteaep["efka_branches"] if b["code"] == "auxiliary")
    assert aux["label"] == "ΕΤΕΑΕΠ"
    assert aux["employee_percent"] == 3.0
    assert eteaep["epikourikiki_kod"] == "001"

    teka = payroll_for_employee(
        employee=_salary_employee(),
        contract={**_salary_contract(), "epikourikiki_kod": "002"},
        params=params,
        period_type="month",
    )
    teka_br = next(b for b in teka["efka_branches"] if b["code"] == "teka")
    assert teka_br["label"] == "ΤΕΚΑ"
    assert teka_br["employee_percent"] == 3.0
    assert not any(b["code"] == "auxiliary" for b in teka["efka_branches"])
    assert teka["efka_employee"] == eteaep["efka_employee"]
    assert eteaep["auxiliary_fund_label"] == "ΕΤΕΑΕΠ"
    assert teka["auxiliary_fund_label"] == "ΤΕΚΑ"


def test_teka_rates_live_in_efka_catalog_group():
    teka = next(item for item in PARAMETER_CATALOG if item["code"] == "efka_teka_employee_percent")
    assert teka["group_code"] == "efka"
    assert teka["group_label"] == "Εισφορές ΕΦΚΑ"


def test_teka_rates_are_separate_from_eteaep():
    params = default_parameter_map()
    params["efka_teka_employee_percent"] = "2"
    params["efka_teka_employer_percent"] = "2"
    eteaep = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=params,
        period_type="month",
    )
    teka = payroll_for_employee(
        employee=_salary_employee(),
        contract={**_salary_contract(), "epikourikiki_kod": "002-ΤΕΚΑ"},
        params=params,
        period_type="month",
    )
    assert teka["efka_employee"] < eteaep["efka_employee"]
    teka_br = next(b for b in teka["efka_branches"] if b["code"] == "teka")
    assert teka_br["employee_percent"] == 2.0


def test_occupational_fund_without_rates_is_zero_and_warns():
    params = default_parameter_map()
    row = payroll_for_employee(
        employee=_salary_employee(),
        contract={**_salary_contract(), "epikourikiki_kod": "004"},
        params=params,
        period_type="month",
    )
    occ = next(b for b in row["efka_branches"] if b["code"] == "aux_004")
    assert occ["label"] == "ΤΕΑΥΦΕ"
    assert occ["employee_percent"] == 0.0
    assert occ["employer_percent"] == 0.0
    assert any("ΤΕΑΥΦΕ" in w for w in row["warnings"])
    eteaep = payroll_for_employee(
        employee=_salary_employee(),
        contract=_salary_contract(),
        params=params,
        period_type="month",
    )
    assert row["efka_employee"] < eteaep["efka_employee"]


def test_payroll_export_xlsx_has_core_sheets():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.payroll_export import build_payroll_export_xlsx

    params = default_parameter_map()
    report = build_payroll_report(
        {"employees": [_salary_employee()]},
        {"111111111": {**_salary_contract(), "epikourikiki_kod": "002"}},
        params,
        period_type="month",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
    )
    content = build_payroll_export_xlsx(
        report=report,
        store={"name": "ΕΡΑΤΩ", "employer_afm": "123456789", "branch_aa": "0"},
        meta_line="ΕΡΑΤΩ · 01/09/2026–30/09/2026",
    )
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == ["Σύνοψη", "Γραμμές", "ΕΦΚΑ", "ΦΜΥ", "Στοιχεία"]
    summary = wb["Σύνοψη"]
    assert summary["A4"].value == "ΔΟΚΙΜΗ Α"
    assert summary["D4"].value == "ΤΕΚΑ"
    totals = [cell.value for cell in summary[summary.max_row]]
    assert totals[0] == "Σύνολο"
    efka = wb["ΕΦΚΑ"]
    labels = [row[4].value for row in efka.iter_rows(min_row=4, max_col=5, values_only=False)]
    assert "ΤΕΚΑ" in labels


def test_apd_preview_flags_missing_identity():
    from app.payroll import attach_apd_identity, build_payroll_report

    params = default_parameter_map()
    timekeeping = {
        "employees": [_salary_employee()],
        "days": [
            {
                "employee_afm": "111111111",
                "work_date": "2026-09-01",
                "recognized_work_minutes": 480,
            },
            {
                "employee_afm": "111111111",
                "work_date": "2026-09-02",
                "recognized_work_minutes": 0,
            },
            {
                "employee_afm": "111111111",
                "work_date": "2026-09-03",
                "recognized_work_minutes": 240,
            },
        ],
    }
    report = build_payroll_report(
        timekeeping,
        {"111111111": _salary_contract()},
        params,
        period_type="month",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
    )
    attach_apd_identity(
        report,
        store={"ame": "", "employer_afm": "123456789", "kad_code": "56.10"},
        employees_by_afm={"111111111": {"amka": "", "amika": ""}},
        timekeeping=timekeeping,
    )
    apd = report["employees"][0]["apd"]
    assert apd["ready"] is False
    assert apd["insurance_days"] == 2
    assert apd["employer_kad"] == "56.10"
    assert any("ΑΜΕ" in gap for gap in apd["gaps"])
    assert any("ΑΜΚΑ" in gap for gap in apd["gaps"])
    assert any("ΑΜΑ" in gap for gap in apd["gaps"])


def test_apd_preview_ready_when_ids_present():
    from app.payroll import attach_apd_identity, build_payroll_report

    params = default_parameter_map()
    report = build_payroll_report(
        {"employees": [_salary_employee()]},
        {"111111111": _salary_contract()},
        params,
        period_type="month",
        period_from=date(2026, 9, 1),
        period_to=date(2026, 9, 30),
    )
    attach_apd_identity(
        report,
        store={"ame": "1234567890", "employer_afm": "123456789", "kad_code": "56.10"},
        employees_by_afm={"111111111": {"amka": "15039012345", "amika": "987654321"}},
        timekeeping={"days": []},
    )
    apd = report["employees"][0]["apd"]
    assert apd["ready"] is True
    assert apd["gaps"] == []
    assert apd["employer_ame"] == "1234567890"
    assert apd["amka"] == "15039012345"
    assert apd["ama"] == "987654321"
    assert apd["packages"]
    assert 'AMKA="15039012345"' in apd["xml"]
    assert report["apd_xml"].startswith("<?xml")
    assert 'AME="1234567890"' in report["apd_xml"]


def test_partial_month_salary_uses_pay_days_and_preserves_full_bonus_base():
    contract = _salary_contract(hire_date="2026-12-21")
    row = payroll_for_employee(employee=_salary_employee(), contract=contract,
        params=default_parameter_map(), period_from=date(2026, 12, 1), period_to=date(2026, 12, 31),
        schedule_rows=_declared_schedule(2026,12))
    # Nine Mon-Fri days, including paid holidays: 9 * 6/5 = 10.8 twenty-fifths.
    assert row["salary_payable_days"] == 10.8
    assert row["period_salary"] == 432
    assert row["salary_full_period"] == 1000
    assert row["bonuses_total"] > 0
    from app.payroll import seasonal_bonus_lines_for
    bonus = seasonal_bonus_lines_for(contract=contract, params=default_parameter_map(),
        period_type="month", period_from=date(2026,12,1), period_to=date(2026,12,31),
        hourly=Decimal(str(row["hourly_wage"])), period_salary=Decimal(1000),
        allowances=Decimal(0), warnings=[])
    assert row["bonuses_total"] == sum(item["amount"] for item in bonus)


def test_salary_departure_fixed_end_and_no_overlap():
    from app.payroll import period_salary_amount
    params = default_parameter_map()
    def amount(**extra):
        return period_salary_amount(_salary_contract(**extra), params, period_type="month",
            period_from=date(2026,9,1), period_to=date(2026,9,30),
            schedule_rows=_declared_schedule(2026,9,six_day=extra.get("weekly_work_days")=="6"))
    assert amount(departure_date="2026-09-10") == Decimal("384.00")
    assert amount(fixed_term_to="2026-09-10", departure_date="2026-09-20") == Decimal("384.00")
    assert amount(hire_date="2026-10-01") == 0
    assert amount(departure_date="2026-08-31") == 0
    assert amount(hire_date="2026-09-20", weekly_work_days="6") == Decimal("360.00")
    assert amount(hire_date="2026-09-01") == 1000


def test_full_february_salary_not_reduced_and_worker_unchanged():
    from app.payroll import period_salary_amount
    for year in (2024, 2026):
        import calendar
        assert period_salary_amount(_salary_contract(hire_date=f"{year}-02-01"),
            default_parameter_map(), period_type="month", period_from=date(year,2,1),
            period_to=date(year,2,calendar.monthrange(year,2)[1])) == 1000
    assert period_salary_amount(_salary_contract(characterization="0"),
        default_parameter_map(), period_type="month", period_from=date(2026,9,1),
        period_to=date(2026,9,30)) == 0
