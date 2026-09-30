from copy import deepcopy
from datetime import date
from io import BytesIO

import pytest
from flask import Flask
from openpyxl import load_workbook

from app.payroll import build_payroll_report, default_parameter_map, attach_apd_identity
from app.payroll_adjustments import apply_payroll_adjustments
from app.payroll_export import build_payroll_export_xlsx


def report():
    return build_payroll_report(
        {"employees": [{"employee_afm": "111111111", "premium_minutes": {"night": 120}}]},
        {"111111111": {"salary": "1000", "weekly_hours": "40", "characterization": "1"}},
        default_parameter_map(), period_from=date(2026, 9, 1), period_to=date(2026, 9, 30),
    )


def test_edits_recompute_totals_and_keep_manual_withholdings_in_excel_and_apd():
    original = report()
    edit = deepcopy(original["employees"][0])
    edit.update(period_salary=500, fmy=12, fmy_manual=True,
                efka_insurable=450, efka_insurable_manual=True,
                total=999999, net=999999, amka="spoofed")
    edit["lines"][0]["hours"] = 4
    edit["efka_branches"][0]["employee_percent"] = 10
    adjusted = apply_payroll_adjustments(original, [edit])
    row = adjusted["employees"][0]
    assert row["period_salary"] == 500
    assert row["total"] != 999999
    assert row["efka_insurable"] == 450
    assert row["efka_branches"][0]["employee_amount"] == 45
    assert row["fmy"] == 12
    assert row["net"] == round(row["after_efka"] - 12, 2)
    assert adjusted["grand_net"] == row["net"]
    assert original["employees"][0]["period_salary"] == 1000
    attach_apd_identity(adjusted, store={"ame": "123"}, employees_by_afm={
        "111111111": {"amka": "12345678901", "amika": "123"}})
    assert row["apd"]["insurable"] == 450
    assert "spoofed" not in adjusted["apd_xml"]
    assert 'asfalismenou="45.00"' in adjusted["apd_xml"]
    wb = load_workbook(BytesIO(build_payroll_export_xlsx(report=adjusted, store={}, meta_line="test")))
    assert any(cell.value == row["net"] for ws in wb for cells in ws for cell in cells)


@pytest.mark.parametrize("value", [-1, "NaN", "Infinity", "oops", 100000001])
def test_invalid_numeric_edits_rejected(value):
    with pytest.raises(ValueError):
        apply_payroll_adjustments(report(), [{"employee_afm": "111111111", "period_salary": value}])


def test_unknown_duplicate_and_changed_lines_rejected():
    original = report()
    with pytest.raises(ValueError):
        apply_payroll_adjustments(original, [{"employee_afm": "222222222"}])
    with pytest.raises(ValueError):
        apply_payroll_adjustments(original, [{"employee_afm": "111111111"}] * 2)
    edit = deepcopy(original["employees"][0])
    edit["lines"][0]["line_kind"] = "bonus"
    with pytest.raises(ValueError):
        apply_payroll_adjustments(original, [edit])


def test_export_route_applies_edits_before_apd_and_excel(monkeypatch):
    from app import routes_payroll as routes
    app = Flask(__name__)
    app.register_blueprint(routes.payroll_bp)
    context = {"id": 1, "name": "Test", "employer_afm": "123456789"}
    monkeypatch.setattr(routes, "resolve_active_store", lambda: context)
    monkeypatch.setattr(routes, "tables_available", lambda: True)
    monkeypatch.setattr(routes, "ensure_seeded", lambda: None)
    monkeypatch.setattr(routes, "list_history_for_store", lambda *a, **kw: [])
    monkeypatch.setattr(routes, "list_schedule_for_range", lambda *a, **kw: [])
    monkeypatch.setattr(routes, "_build_timekeeping_for_month", lambda *a, **kw: (
        {"employees": [{"employee_afm": "111111111"}]}, [], {}))
    monkeypatch.setattr(routes, "_contracts_by_afm", lambda ctx: {
        "111111111": {"salary": "1000", "weekly_hours": "40", "characterization": "1"}})
    monkeypatch.setattr(routes, "load_resolved", lambda **kw: default_parameter_map())
    monkeypatch.setattr(routes, "list_schedule_for_range", lambda *a, **kw: [])
    monkeypatch.setattr(routes, "list_employees_for_employer", lambda *a, **kw: [])
    monkeypatch.setattr(routes.repo_store, "get_store_config", lambda sid: context)
    client = app.test_client()
    payload = {"year": 2026, "month": 9, "store_id": 1,
               "adjustments": [{"employee_afm": "111111111", "period_salary": 500}]}
    response = client.post("/api/payroll/export", json={**payload, "format": "apd-preview"})
    assert response.status_code == 200
    data = response.get_json()
    assert data["employees"][0]["total"] == 500
    assert data["employees"][0]["apd"]["insurable"] == 500
    assert "500.00" in data["apd_xml"]
    response = client.post("/api/payroll/export", json=payload)
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.data))
    assert any(cell.value == 500 for ws in wb for cells in ws for cell in cells)
    assert client.post("/api/payroll/export", json={**payload, "store_id": 2}).status_code == 400
    payload["adjustments"][0]["period_salary"] = "NaN"
    assert client.post("/api/payroll/export", json=payload).status_code == 400


def test_browser_day_edit_and_exports_share_current_values(tmp_path):
    import json
    import shutil
    import subprocess
    from pathlib import Path
    node = shutil.which("node")
    if not node:
        pytest.skip("Node required for browser calculation regression")
    data = report()
    data["store"] = {"id": 1}
    script = """
const fs = require('fs');
const vm = require('vm');
const assert = require('node:assert/strict');
const inputs = {};
const context = vm.createContext({URLSearchParams, Intl, location:{search:'?year=2026&month=9'},
  document:{addEventListener(){}, querySelector(){return null}, querySelectorAll(){return []}, getElementById(){return null}},
  Office:{escapeHtml:s=>s}});
vm.runInContext(fs.readFileSync('app/static/js/apologistic-timekeeping.js','utf8'),context);
vm.runInContext('payrollData = '+fs.readFileSync(process.argv[2],'utf8')+'; payrollOriginal=JSON.parse(JSON.stringify(payrollData)); payrollModalIndex=0; renderPayroll=()=>{};',context);
vm.runInContext(`onPayrollModalInput({target:{closest(){return {value:'2',getAttribute(k){return k==='data-emp'?'salary_unpaid_days':null}}}}});`, context);
const value = vm.runInContext('payrollExportPayload()',context);
assert.equal(value.adjustments[0].period_salary,920);
assert.equal(value.adjustments[0].salary_unpaid_days,2);
assert.equal(value.store_id,1);
"""
    fixture = tmp_path / "report.json"
    fixture.write_text(json.dumps(data), encoding="utf-8")
    runner = tmp_path / "check.cjs"
    runner.write_text(script, encoding="utf-8")
    subprocess.run([node, str(runner), str(fixture)], check=True,
        cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True)
