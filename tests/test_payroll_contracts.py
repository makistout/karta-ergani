from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.payroll import build_payroll_report, default_parameter_map, payroll_for_employee
from app.payroll_contracts import contract_segments
from app.payroll_adjustments import apply_payroll_adjustments
from app.payroll_export import build_payroll_export_xlsx


def contract(effective, salary='1000', **kw):
    return dict(employee_afm='111111111', ergani_updated_at=effective, salary=salary,
                weekly_hours='40', characterization='1', hourly_wage='10', **kw)


def calculate(history, start=date(2026, 9, 1), end=date(2026, 9, 30), period='month', parameters=None):
    days, schedule = [], []
    when = start
    while when <= end:
        if when.weekday() < 5:
            days.append(dict(employee_afm='111111111', work_date=when.isoformat(),
                             premium_minutes={'night': 60}, overtime_40_breakdown={'day': 60}))
            schedule.append(dict(employee_afm='111111111', work_date=when.isoformat(),
                                 hour_from='09:00', hour_to='17:00'))
        when += timedelta(days=1)
    employee = dict(employee_afm='111111111', premium_minutes={'night': len(days)*60},
                    overtime_40_breakdown={'day': len(days)*60})
    params = default_parameter_map()
    params['zone_premium_base'] = 'contractual'
    params.update(parameters or {})
    return build_payroll_report(dict(employees=[employee], days=days),
        {'111111111': history[-1]}, params, period_type=period,
        period_from=start, period_to=end, schedule_rows=schedule, contract_history=history)


def test_no_change_keeps_existing_calculation_and_format_updates_coalesce():
    first = contract('2026-01-01', '1.000,00', weekly_work_days='5')
    last = contract('2026-09-16', '1000', weekly_work_days='5-ήμερη')
    assert len(contract_segments([first, last], last, date(2026,9,1), date(2026,9,30))) == 1
    result = calculate([first, last])['employees'][0]
    baseline = calculate([first])['employees'][0]
    assert result['total'] == baseline['total']
    assert 'contract_segments' not in result


def test_month_three_contracts_split_at_inclusive_effective_dates():
    history = [contract('2026-01-01', '1000'), contract('2026-09-11', '1200'), contract('2026-09-21', '1500')]
    row = calculate(history)['employees'][0]
    segments = row['contract_segments']
    assert [(s['from'], s['to']) for s in segments] == [
        ('2026-09-01','2026-09-10'), ('2026-09-11','2026-09-20'), ('2026-09-21','2026-09-30')]
    # 8, 6 and 8 declared workdays, normalized by 40 contractual hours/week.
    assert [s['period_salary'] for s in segments] == [363.64,327.27,545.45]
    assert row['period_salary'] == 1236.36
    assert sum(l['hours'] for l in row['lines'] if l['family'] == 'Βάση') == 22
    assert row['total'] == round(row['period_salary'] + sum(l['amount'] for l in row['lines']),2)


def test_week_change_uses_each_rate_and_tax_once():
    old = contract('2026-01-01', '1000')
    new = contract('2026-09-16', '1500')
    new['hourly_wage'] = '15'
    result = calculate([old,new], date(2026,9,14), date(2026,9,20), 'week')
    row = result['employees'][0]
    assert [s['salary_share'] for s in row['contract_segments']] == [0.4,0.6]
    assert row['period_salary'] == 312.06
    overtime = [l for l in row['lines'] if l['family'] != 'Βάση' and l['line_kind'] == 'hour']
    assert [(l['hours'], l['hourly']) for l in overtime] == [(2,10),(3,15)]
    from app.payroll import recalculate_employee_row
    assert recalculate_employee_row(row,params=result['parameters'])['net'] == row['net']


def test_future_contract_not_applied_to_old_period():
    row = calculate([contract('2026-01-01','1000'), contract('2026-10-01','1500')])['employees'][0]
    assert row['period_salary'] == 1000
    assert 'contract_segments' not in row


def test_same_effective_date_latest_snapshot_wins_and_boundary_has_no_empty_segment():
    old = contract('2026-01-01','1000', id=1)
    a = contract('2026-09-01','1200', id=2)
    b = contract('2026-09-01','1500', id=3)
    segments = contract_segments([b,old,a], b, date(2026,9,1),date(2026,9,30))
    assert len(segments) == 1
    assert segments[0]['contract']['salary'] == '1500'


def test_workers_split_hourly_rates_without_requiring_declared_schedule():
    from app.payroll_contracts import segmented_payroll
    first, last = contract('2026-01-01'), contract('2026-09-16')
    first.update(characterization='ΕΡΓΑΤΗΣ')
    last.update(characterization='ΕΡΓΑΤΗΣ', hourly_wage='15')
    start,end = date(2026,9,14),date(2026,9,20)
    days = [dict(work_date='2026-09-15',premium_minutes={'day':60}),dict(work_date='2026-09-16',premium_minutes={'day':60})]
    row = segmented_payroll({'premium_minutes':{'day':120}}, contract_segments([first,last],last,start,end),days,None,default_parameter_map(),'week',start,end)
    assert row['period_salary'] == 0
    assert row['base_pay'] == 25


def test_missing_daily_detail_rejected():
    from app.payroll_contracts import segmented_payroll
    first,last = contract('2026-01-01'),contract('2026-09-16','1500')
    start,end=date(2026,9,1),date(2026,9,30)
    with pytest.raises(ValueError,match='ημερήσια'):
        segmented_payroll({},contract_segments([first,last],last,start,end),[],[],default_parameter_map(),'month',start,end)


def test_exports_keep_segment_identity_and_reject_flattening():
    report = calculate([contract('2026-01-01'),contract('2026-09-16','1500')])
    edit=deepcopy(report['employees'][0])
    assert apply_payroll_adjustments(report,[edit])['grand_total'] == report['grand_total']
    edit['lines'][0]['contract_segment']='wrong'
    with pytest.raises(ValueError,match='γραμμές'):
        apply_payroll_adjustments(report,[edit])
    edit=deepcopy(report['employees'][0])
    edit['hourly_wage']=99
    with pytest.raises(ValueError,match='επιμερισμό'):
        apply_payroll_adjustments(report,[edit])
    wb=load_workbook(BytesIO(build_payroll_export_xlsx(report=report,store={},meta_line='test')))
    assert 'Συμβάσεις' in wb.sheetnames
    assert any('01/09/2026' in str(cell.value) for cells in wb['Γραμμές'] for cell in cells)


def test_monthly_bonus_is_split_and_lump_sum_is_not_duplicated():
    history=[contract('2026-01-01'),contract('2026-09-16','1500')]
    monthly=calculate(history,parameters={'bonus_payout':'κάθε_μήνας_περιόδου'})['employees'][0]
    gifts=[l for l in monthly['lines'] if l['line_kind']=='bonus']
    assert len(gifts)==2
    assert {l['contract_segment'] for l in gifts} == {'01/09/2026–15/09/2026','16/09/2026–30/09/2026'}
    lump=calculate(history,parameters={'bonus_christmas_pay_month':'9'})['employees'][0]
    assert len([l for l in lump['lines'] if l['line_kind']=='bonus'])==1


def test_contract_hours_affect_salary_shares():
    first,last=contract('2026-01-01'),contract('2026-09-16','1500')
    first['weekly_hours']='20'
    row=calculate([first,last])['employees'][0]
    assert row['contract_segments'][0]['salary_share'] > 0.5
    assert row['period_salary'] < 1250


def test_missing_or_mismatched_daily_minutes_cannot_silently_drop_pay():
    from app.payroll_contracts import segmented_payroll
    first,last=contract('2026-01-01'),contract('2026-09-16','1500')
    start,end=date(2026,9,1),date(2026,9,30)
    with pytest.raises(ValueError,match='δεν συμφωνεί'):
        segmented_payroll({'night':120},contract_segments([first,last],last,start,end),
            [{'work_date':'2026-09-15','premium_minutes':{'night':60}}],[],default_parameter_map(),'month',start,end)


def test_browser_preserves_segment_amounts_and_displays_intervals(tmp_path):
    import json, shutil, subprocess
    from pathlib import Path
    node=shutil.which('node')
    if not node: pytest.skip('Node required')
    report=calculate([contract('2026-01-01'),contract('2026-09-16','1500')])
    fixture=tmp_path/'report.json'
    fixture.write_text(json.dumps(report),encoding='utf8')
    script=tmp_path/'check.cjs'
    script.write_text("""
const fs=require('fs'),vm=require('vm'),assert=require('node:assert/strict');
const elements={},inputs={};
const context=vm.createContext({URLSearchParams,Intl,location:{search:'?year=2026&month=9'},
 document:{addEventListener(){},getElementById(k){return elements[k]??={}},querySelector(k){return inputs[k]??={}},querySelectorAll(){return []}},
 Office:{escapeHtml:s=>s}});
vm.runInContext(fs.readFileSync('app/static/js/apologistic-timekeeping.js','utf8'),context);
vm.runInContext('payrollData='+fs.readFileSync(process.argv[2],'utf8')+'; payrollModalIndex=0;',context);
const before=vm.runInContext('payrollData.employees[0].total',context);
vm.runInContext('applyEmployee(payrollData.employees[0]);fillPayrollModal(payrollData.employees[0]);',context);
assert.equal(vm.runInContext('payrollData.employees[0].total',context),before);
assert.ok(elements.payrollInfoModalBody.innerHTML.includes('2026-09-16'));
assert.ok(inputs['[data-emp="hourly_wage"]'].disabled);
vm.runInContext(`onPayrollModalInput({target:{closest(){return {value:'99',getAttribute(k){return k==='data-emp'?'hourly_wage':null}}}}});`,context);
assert.equal(vm.runInContext('payrollData.employees[0].hourly_wage',context),10);
""",encoding='utf8')
    subprocess.run([node,str(script),str(fixture)],check=True,cwd=Path(__file__).resolve().parents[1],capture_output=True,text=True)
