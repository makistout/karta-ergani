"""Excel εξαγωγή υπολογισμένης μισθοδοσίας."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.timekeeping_export import _style_sheet

_MONEY_FMT = "#,##0.00"
_PCT_FMT = "0.00"
_BORDER = Side(style="thin", color="D9E2E9")
_FILL_LIGHT = PatternFill("solid", fgColor="EAF2F8")
_FONT_BODY = Font(name="Aptos", size=10)
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="top")
_ALIGN_TOP = Alignment(vertical="top")


def _name(row: dict[str, Any]) -> str:
    return f"{row.get('eponymo') or ''} {row.get('onoma') or ''}".strip()


def _num(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _aux_branch(row: dict[str, Any]) -> dict[str, Any]:
    for branch in row.get("efka_branches") or []:
        code = str(branch.get("code") or "")
        if code in {"teka", "auxiliary"} or code.startswith("aux_"):
            return branch
    return {}


def _date_el(value: Any) -> str:
    text = str(value or "")[:10]
    if len(text) == 10 and text[4] == "-" and text[7] == "-":
        return f"{text[8:10]}/{text[5:7]}/{text[:4]}"
    return text


def _extras(row: dict[str, Any]) -> float:
    return round(_num(row.get("base_pay")) + _num(row.get("extra_pay")), 2)


def _finish_money(ws, header_row: int, money_cols: set[int], pct_cols: set[int] | None = None) -> None:
    max_col = ws.max_column
    data_rows = max(0, ws.max_row - header_row)
    if data_rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{ws.max_row}"
    pct_cols = pct_cols or set()
    for row in range(header_row + 1, ws.max_row + 1):
        zebra = row % 2 == 0
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            cell.font = _FONT_BODY
            cell.border = Border(bottom=_BORDER)
            if zebra:
                cell.fill = _FILL_LIGHT
            if col in money_cols:
                cell.number_format = _MONEY_FMT
                cell.alignment = _ALIGN_RIGHT
            elif col in pct_cols:
                cell.number_format = _PCT_FMT
                cell.alignment = _ALIGN_RIGHT
            else:
                cell.alignment = _ALIGN_TOP


def build_payroll_export_xlsx(
    *,
    report: dict[str, Any],
    store: dict[str, Any],
    meta_line: str,
) -> bytes:
    employees = list(report.get("employees") or [])
    wb = Workbook()

    summary = wb.active
    summary.title = "Σύνοψη"
    summary_headers = [
        "Εργαζόμενος", "ΑΦΜ", "Χαρακτηρισμός", "Επικουρική",
        "Ημ/νία πρόσληψης", "Μισθός περιόδου", "Προσαυξήσεις", "Επιδόματα",
        "Δώρα / άδεια", "Μεικτά", "ΕΦΚΑ ασφαλισμένου", "Επικουρική €",
        "ΦΜΥ", "Καθαρά", "ΕΦΚΑ εργοδότη", "Παρατηρήσεις",
    ]
    header = _style_sheet(
        summary,
        title="Μισθοδοσία",
        meta=meta_line,
        headers=summary_headers,
        widths=[28, 12, 14, 16, 14, 14, 14, 12, 13, 12, 16, 13, 12, 12, 14, 40],
    )
    for row in employees:
        aux = _aux_branch(row)
        summary.append([
            _name(row),
            str(row.get("employee_afm") or ""),
            row.get("characterization_label") or "",
            row.get("auxiliary_fund_label") or aux.get("label") or "",
            _date_el(row.get("hire_date")),
            _num(row.get("period_salary")),
            _extras(row),
            _num(row.get("allowances_total")),
            _num(row.get("bonuses_total")),
            _num(row.get("total")),
            _num(row.get("efka_employee")),
            _num(aux.get("employee_amount")),
            _num(row.get("fmy")),
            _num(row.get("net")),
            _num(row.get("efka_employer")),
            " · ".join(str(item) for item in (row.get("warnings") or [])),
        ])
    _finish_money(summary, header, set(range(6, 16)))
    if employees:
        total_row = summary.max_row + 1
        summary.append([
            "Σύνολο", "", "", "", "",
            sum(_num(r.get("period_salary")) for r in employees),
            sum(_extras(r) for r in employees),
            sum(_num(r.get("allowances_total")) for r in employees),
            sum(_num(r.get("bonuses_total")) for r in employees),
            _num(report.get("grand_total")),
            _num(report.get("grand_efka_employee")),
            sum(_num(_aux_branch(r).get("employee_amount")) for r in employees),
            _num(report.get("grand_fmy")),
            _num(report.get("grand_net")),
            _num(report.get("grand_efka_employer")),
            "",
        ])
        for col in range(1, 17):
            cell = summary.cell(total_row, col)
            cell.font = Font(name="Aptos", size=10, bold=True)
            if col >= 6:
                cell.number_format = _MONEY_FMT
                cell.alignment = _ALIGN_RIGHT

    lines_ws = wb.create_sheet("Γραμμές")
    line_headers = [
        "Εργαζόμενος", "ΑΦΜ", "Είδος", "Οικογένεια", "Ζώνη", "Ώρες",
        "% οικογένειας", "% ζώνης", "Υπολογισμός", "Ποσό",
    ]
    line_header = _style_sheet(
        lines_ws,
        title="Ανάλυση γραμμών μισθοδοσίας",
        meta=meta_line,
        headers=line_headers,
        widths=[28, 12, 12, 22, 16, 10, 14, 12, 48, 12],
    )
    kind_label = {"hour": "Ώρα", "allowance": "Επίδομα", "bonus": "Δώρο/άδεια"}
    for row in employees:
        for line in row.get("lines") or []:
            kind = str(line.get("line_kind") or "hour")
            lines_ws.append([
                _name(row),
                str(row.get("employee_afm") or ""),
                kind_label.get(kind, kind),
                line.get("family") or "",
                line.get("zone") or "",
                _num(line.get("hours")),
                _num(line.get("family_percent")),
                _num(line.get("zone_percent")),
                line.get("formula") or "",
                _num(line.get("amount")),
            ])
    _finish_money(lines_ws, line_header, {6, 10}, {7, 8})

    efka_ws = wb.create_sheet("ΕΦΚΑ")
    efka_headers = [
        "Εργαζόμενος", "ΑΦΜ", "Επικουρική", "Ασφαλιστέες αποδοχές",
        "Κλάδος", "% ασφαλισμένου", "Κράτηση €", "% εργοδότη", "Εργοδότης €",
        "Υπολογισμός",
    ]
    efka_header = _style_sheet(
        efka_ws,
        title="Εισφορές ΕΦΚΑ / ΤΕΚΑ",
        meta=meta_line,
        headers=efka_headers,
        widths=[28, 12, 14, 18, 22, 16, 12, 12, 14, 48],
    )
    for row in employees:
        for branch in row.get("efka_branches") or []:
            efka_ws.append([
                _name(row),
                str(row.get("employee_afm") or ""),
                row.get("auxiliary_fund_label") or "",
                _num(row.get("efka_insurable")),
                branch.get("label") or "",
                _num(branch.get("employee_percent")),
                _num(branch.get("employee_amount")),
                _num(branch.get("employer_percent")),
                _num(branch.get("employer_amount")),
                branch.get("formula") or "",
            ])
    _finish_money(efka_ws, efka_header, {4, 7, 9}, {6, 8})

    tax_ws = wb.create_sheet("ΦΜΥ")
    tax_headers = [
        "Εργαζόμενος", "ΑΦΜ", "Ηλικιακή ομάδα", "Ετήσιο φορολογητέο",
        "Μείωση άρθ. 16", "ΦΜΥ περιόδου", "Κλιμάκιο", "%", "Φόρος κλιμακίου",
        "Υπολογισμός",
    ]
    tax_header = _style_sheet(
        tax_ws,
        title="ΦΜΥ / φόρος εισοδήματος",
        meta=meta_line,
        headers=tax_headers,
        widths=[28, 12, 16, 18, 16, 14, 16, 10, 16, 48],
    )
    for row in employees:
        brackets = list(row.get("tax_brackets") or [])
        if not brackets:
            tax_ws.append([
                _name(row),
                str(row.get("employee_afm") or ""),
                row.get("tax_age_group") or "",
                _num(row.get("tax_annual")),
                _num(row.get("tax_credit_annual")),
                _num(row.get("fmy")),
                "", "", "",
                row.get("tax_formula") or "",
            ])
            continue
        for bracket in brackets:
            upto = bracket.get("upto")
            tax_ws.append([
                _name(row),
                str(row.get("employee_afm") or ""),
                row.get("tax_age_group") or "",
                _num(row.get("tax_annual")),
                _num(row.get("tax_credit_annual")),
                _num(row.get("fmy")),
                "Υπερβάλλον" if upto is None else f"Έως {upto}",
                _num(bracket.get("percent")),
                _num(bracket.get("amount")),
                bracket.get("formula") or row.get("tax_formula") or "",
            ])
    _finish_money(tax_ws, tax_header, {4, 5, 6, 9}, {8})

    info = wb.create_sheet("Στοιχεία")
    info_headers = ["Πεδίο", "Τιμή"]
    info_header = _style_sheet(
        info,
        title="Περίοδος / κατάστημα",
        meta=meta_line,
        headers=info_headers,
        widths=[28, 60],
    )
    info.append(["Κατάστημα", store.get("name") or ""])
    info.append(["ΑΦΜ εργοδότη", store.get("employer_afm") or ""])
    info.append(["Παράρτημα", store.get("branch_aa") if store.get("branch_aa") is not None else ""])
    info.append(["Περίοδος", report.get("period_type") or ""])
    info.append(["Από", _date_el(report.get("period_from"))])
    info.append(["Έως", _date_el(report.get("period_to"))])
    info.append(["Έκδοση", report.get("calculation_version") or ""])
    info.append(["Εργαζόμενοι", int(report.get("counts", {}).get("employees") or len(employees))])
    _finish_money(info, info_header, set())

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
